"""Export actual four-level TRAILS graphs, without distributing coefficients.

Run in the TRAILS Python environment. Source helpers and licensed packages stay
outside the public app. These exports run routing only, not LCI/LCIA or FaIR.
Use --case independently to bound memory use and allow resumable generation.
"""

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import pickle
import sys

REPO = Path(__file__).resolve().parents[3]
DATA = REPO / "apps/seoul_2026/data"
OUTPUT = DATA / "routing"
WORK = REPO / "build/seoul-routing"


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def cement_root_workbook(source):
    """Add an explicit one-tonne-stored CCS-minus-equal-clinker-BAU demand."""
    import openpyxl

    basis = json.loads((DATA / "cement_storage_basis.json").read_text())
    amount = 1000 * basis["campaign_clinker_t_per_t_stored"]
    wb = openpyxl.load_workbook(source)
    sheet = wb.active
    database = sheet.cell(1, 2).value
    header = next(list(row) for row in sheet.values if row[0] == "name" and row[1] == "amount")
    identity = dict(name="Seoul | cement CCS minus BAU | 1 tonne CO2 stored",
                    **{"reference product": "CCS minus equal-output clinker BAU, per tonne CO2 stored"}, location="DK")
    rows = [[None], ["Activity", identity["name"]], ["location", "DK"],
            ["production amount", 1], ["reference product", identity["reference product"]],
            ["type", "process"], ["unit", "unit"], ["Exchanges"], header]
    exchanges = [dict(name=identity["name"], amount=1, location="DK", unit="unit", type="production",
                      **{"reference product": identity["reference product"]})]
    for system, sign in (("CCS", 1), ("BAU", -1)):
        exchanges.append(dict(name=f"ccus | clinker cohort 2035-2064 | SSP2-PkBudg1000 | {system}",
            amount=sign*amount, location="DK", unit="kilogram", type="technosphere",
            database=database, **{"reference product": "grey clinker, cohort-normalized"},
            temporal_distribution=1, temporal_loc=0, temporal_amount_source="port"))
    rows.extend([[e.get(c) for c in header] for e in exchanges])
    for row in rows:
        sheet.append(row)
    target = WORK / "cement-normalized-foreground.xlsx"
    wb.save(target)
    return target, identity


def import_cement_foreground(model, path):
    """Preserve unaffected background timing without repeated full-map scans.

    TRAILS' importer clears every affected caller in every template year by
    repeatedly scanning both temporal maps. Partitioning once before importing
    is equivalent: remove existing profiles only for exact foreground identities,
    let the importer create their replacements, then restore untouched callers.
    No library patch or background profile change is needed.
    """
    import openpyxl

    identities = set()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in wb:
        identity = None
        for row in sheet.values:
            if row[0] == "Activity":
                if identity:
                    identities.add(tuple(identity.get(k, "").strip() for k in ("name", "reference product", "location")))
                identity = {"name": str(row[1])}
            elif identity is not None and row[0] in {"reference product", "location"}:
                identity[row[0]] = str(row[1])
        if identity:
            identities.add(tuple(identity.get(k, "").strip() for k in ("name", "reference product", "location")))
    wb.close()
    metadata = next(iter(model.activity_indices.values()))
    overridden = {int(i) for i,m in metadata.items()
        if tuple(str(m.get(k) or "").strip() for k in ("name", "reference product", "location")) in identities}
    saved = {}
    for attribute in ("temporal_technosphere_exchanges", "temporal_biosphere_exchanges"):
        saved[attribute] = {k:v for k,v in getattr(model, attribute).items() if int(k[1]) not in overridden}
        setattr(model, attribute, {})
    print(f"Preserve timing of unaffected background callers; {len(overridden)} exact existing overrides", flush=True)
    model.import_excel_inventory(str(path))
    for attribute, original in saved.items():
        replacement = getattr(model, attribute)
        assert not (original.keys() & replacement.keys())
        original.update(replacement)
        setattr(model, attribute, original)
    return dict(existing_callers_overridden=len(overridden), foreground_identities=len(identities))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", choices=("daccs", "beccs", "erw", "cement"), required=True)
    parser.add_argument("--source-dir", type=Path, default=REPO/"build/setac-regeneration-20261014")
    parser.add_argument("--package", type=Path, default=REPO/"build/seoul-daccs-energy-signs/package/datapackage.json")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    source = args.source_dir.resolve()
    sys.path.insert(0, str(REPO))
    sys.path.insert(0, str(source))
    from scripts import generate_routing_graphs as helpers
    from workshop.temporal import enable_signed_technosphere_routing
    from workshop.lifetime import cohort_cases
    import bw2data as bd
    from datapackage import Package
    import trails
    from trails import Trails
    from apps.seoul_2026.scripts.routing_structure import consolidate_root, boundary_demands, compare_boundaries
    from apps.seoul_2026.scripts.readable_activity_names import historical_carrier_names

    bd.projects.set_current("ecoinvent-3.12-cutoff")
    output_dir = args.output.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    case = args.case
    config_path = source / "data/inventories/parameters.json"
    config = json.loads(config_path.read_text())
    package = args.package.resolve()
    if case == "erw":
        foregrounds = [DATA/"coastal_erw_campaign.xlsx"]
        identity = dict(name="coastal olivine weathering deployment, 20-year campaign, dated reference",
            **{"reference product": "coastal weathering campaign, 1.25 kg rock total"}, location="RER")
        year, amount, floor = 2030, 1000, 1e-15
    elif case == "cement":
        original = REPO/"apps/ccus_webinar/generated/inventories/slide12_diagnostic/a81aabcc1960/foreground_temporal.xlsx"
        fg, identity = cement_root_workbook(original)
        foregrounds = [fg]
        year, amount, floor = 2035, 1, 1e-18
    else:
        name = case.upper()
        foregrounds = [source/f"data/inventories/generated/{f}" for f in
                       ("lcss2026-foreground.xlsx", "lcss2026-routing-cohorts.xlsx")]
        identity = helpers.ROOTS[name]
        year = 2030
        stored = cohort_cases(config, commissioning_year=year)[name].lifetime_net_stored_kg / 1000
        amount = helpers.cohort_root_amount(name, config, year) / stored
        floor = 1e-15 / stored
    record = dict(schema_version=1, case=case, generated_at=datetime.now(timezone.utc).isoformat(),
        trails_version=trails.__version__, trails_source=str(Path(trails.__file__).resolve()),
        scenario="REMIND-EU SSP2-PkBudg1000", ecoinvent="3.12 cut-off", start_year=year,
        routing=dict(max_depth=4, adaptive_relative_score_cutoff=None, min_amount=floor,
                     attribute_to_roots=True),
        functional_unit="1 t modelled eventual gross CO2 uptake potential" if case=="erw" else "1 t CO2 physically stored",
        root_identity={k:identity[k] for k in ("name", "reference product", "location")},
        root_amount=amount, interpolation_end_year_offset=83 if case=="beccs" else 0, foregrounds=[dict(path=str(p), sha256=digest(p)) for p in foregrounds],
        package_descriptor_sha256=digest(package), script_sha256=digest(Path(__file__)),
        scope="Actual temporal_routing graph at fixed depth four. Routing illustration only; no new LCI, LCIA or climate scores. Public export retains topology, activity identities and calendar years, but omits private/background coefficients and amounts.")
    if case == "cement":
        record["source_foreground_sha256"] = digest(original)
        record["comparison"] = "CCS minus equal-output BAU, normalized to 1 t physically stored"
        record["background_vintage_note"] = "October 2026 scenario package for this new routing view; stored cement scores retain their September release package. This is not the original adaptive score graph."
    if case == "erw":
        record["routing_note"] = "Four-level visualization; reviewed scores retain their original fixed depth one."
    print(f"Load {case}: annual background + foreground", flush=True)
    with Trails(Package(str(package)), interpolate_annual=True, cache_interpolation=False,
                interpolation_start_year_offset=0, interpolation_end_year_offset=83 if case=="beccs" else 0, ei_version="3.12") as model:
        # Import synchronizes every new foreground identity across year maps.
        # Share only dictionaries proven exactly equal, avoiding repeated scans
        # of identical background metadata for the large cement foreground.
        shared = 0
        for attribute in ("activity_indices", "biosphere_indices"):
            mappings = getattr(model, attribute)
            canonical = []
            for label, mapping in list(mappings.items()):
                existing = next((m for m in canonical if m == mapping), None)
                if existing is None:
                    canonical.append(mapping)
                else:
                    mappings[label] = existing
                    shared += 1
        print(f"Reuse {shared} exactly equal metadata maps; import foreground", flush=True)
        record["identical_metadata_maps_shared"] = shared
        if case == "cement":
            record["import_partition_checks"] = import_cement_foreground(model, foregrounds[0])
        else:
            model.import_excel_inventory([str(p) for p in foregrounds])
        diagnostics = model.import_diagnostics
        assert not diagnostics.get("unlinked") and not diagnostics.get("new_flows"), "Unlinked foreground"
        if case == "beccs":
            forest = config["beccs"]["forest_feedstock"]
            helpers.reset_background_temporal_exchanges(model, preserve_activity_names={forest["name"], forest["regrowth_name"]})
            helpers.apply_beccs_cohort_timing(model, year, config)
        elif case == "daccs":
            helpers.apply_daccs_cohort_timing(model, year, config)
        enable_signed_technosphere_routing(model)
        root = helpers.find_activity(model, year, identity)
        inlined = set()
        integrated_name = None
        if case == "beccs":
            capture_name = config["beccs"]["capture_module"]["generated_name"]
            inlined.add(helpers.find_activity(model,year,dict(name=capture_name,
                **{"reference product":"carbon dioxide, captured"},location=config["foreground_location"])))
            integrated_name = capture_name
        elif case == "daccs":
            daccs = config["daccs"]
            inlined.add(helpers.find_activity(model,year,dict(name=daccs["wrapper_name"],
                **{"reference product":daccs["reference_product"]},location=config["foreground_location"])))
            native = daccs["native_activity"]
            inlined.add(helpers.find_activity(model,year,dict(name=native["name"],
                **{"reference product":native["reference_product"]},location=native["scenario_location"])))
            integrated_name = "carbon dioxide, captured and stored, solid-sorbent direct air capture with heat pump"
        elif case == "cement":
            from apps.ccus_webinar.lca_model.pipeline import root_identity
            inlined.update(helpers.find_activity(model,year,root_identity("SSP2-PkBudg1000",s,cohort=True)) for s in ("BAU","CCS"))
            integrated_name = "carbon dioxide, captured and stored, at clinker production, relative to matching production without capture"
        if integrated_name:
            print(f"Route original {case} for dated boundary-equivalence checks",flush=True)
            model.temporal_routing(start_year=year,start_act_idx=root,amount=amount,max_depth=4,
                adaptive_relative_score_cutoff=None,min_amount=floor,attribute_to_roots=True,show_progress=False)
            reference = boundary_demands(model.graph,root,inlined)
            record["original_root_identity"] = record["root_identity"]
            record["original_root_amount"] = amount
            identity,structure = consolidate_root(model,root=root,year=year,amount=amount,
                inlined=inlined,name=integrated_name,helpers=helpers)
            record["root_identity"] = identity
            amount = 1000
            record["root_amount"] = amount
            record["structure_review"] = structure
            record["structure_review"]["compiler_sha256"] = digest(REPO/"apps/seoul_2026/scripts/routing_structure.py")
            record["structure_review"]["scope"] = "Equivalent integrated foreground root; retained dated supplier demands and direct biosphere profiles. Existing numerical LCA/FaIR artifacts are unchanged."
        print(f"Route {case}, level 0 to 4", flush=True)
        model.temporal_routing(start_year=year, start_act_idx=root, amount=amount,
            max_depth=4, adaptive_relative_score_cutoff=None, min_amount=floor,
            attribute_to_roots=True, show_progress=False)
        graph = model.graph
        if integrated_name:
            record["structure_review"].update(compare_boundaries(reference,boundary_demands(graph,root,set())))
        with (WORK/f"{case}-integrated-private-graph.pkl").open("wb") as stream:
            pickle.dump(graph, stream)
        ids = {n:i for i,n in enumerate(graph.nodes)}
        readable_names = historical_carrier_names(original) if case == "cement" else {}
        if readable_names:
            record["display_name_postprocessing_sha256"] = digest(REPO/"apps/seoul_2026/scripts/readable_activity_names.py")
        nodes = []
        for n,d in graph.nodes(data=True):
            nodes.append(dict(id=ids[n], depth=int(d["depth"]), year=int(d["year"]),
                name=str(d.get("name", "")), product=str(d.get("reference_product", "")),
                location=str(d.get("location", "")), unit=str(d.get("unit", ""))))
            if nodes[-1]["name"] in readable_names:
                nodes[-1]["display_name"] = readable_names[nodes[-1]["name"]]
        edges = [dict(source=ids[u], target=ids[v], negative=float(d.get("amount",0)) < 0)
                 for u,v,d in graph.edges(data=True)]
        assert min(n["depth"] for n in nodes) == 0 and max(n["depth"] for n in nodes) <= 4
        assert sum(n["depth"] == 0 for n in nodes) == 1
        record.update(nodes=len(nodes), edges=len(edges), years=[min(n["year"] for n in nodes),max(n["year"] for n in nodes)],
                      depth_counts={str(i):sum(n["depth"]==i for n in nodes) for i in range(5)})
        output = output_dir/f"{case}.json"
        output.write_text(json.dumps(dict(metadata=record,nodes=nodes,edges=edges),separators=(",", ":"))+"\n")
        print(f"Saved {output.name}: {len(nodes)} nodes / {len(edges)} edges; SHA {digest(output)}", flush=True)


if __name__ == "__main__":
    main()
